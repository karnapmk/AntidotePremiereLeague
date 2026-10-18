import logging
import re
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib import font_manager
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.offsetbox import AnnotationBbox, OffsetImage
from matplotlib.patches import Circle, Polygon, Rectangle, RegularPolygon, Wedge
from matplotlib.ticker import FuncFormatter, MaxNLocator
from matplotlib.transforms import Affine2D
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PIL import Image

matplotlib.rcParams["pdf.fonttype"] = 42
# macOS .ttc faces can't be fully subset; the notices are harmless but noisy.
logging.getLogger("fontTools").setLevel(logging.ERROR)

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "AntidoteScores.txt"
WORKBOOK = BASE_DIR / "AntidoteStats.xlsx"
REPORT_PDF = BASE_DIR / "Antidote Report Card.pdf"
# Badge art lives on disk so the script needs no font tooling at runtime. The two
# emoji were pulled from Apple Color Emoji's precomposed glyphs (u1F4AA.4 is the
# brown muscle; writing the skin-tone modifier as text needs a shaper PIL lacks).
IMAGE_DIR = BASE_DIR / "Images"
SURVIVOR_LOGO = IMAGE_DIR / "Survivor_US_TV_logo.png"
MUSCLE_BADGE = IMAGE_DIR / "badge_muscle_brown.png"
TROPHY_BADGE = IMAGE_DIR / "badge_trophy.png"
# Keyed off their source JPEGs: the supplied art had a white backdrop and a painted
# transparency checkerboard, both of which would show as boxes on the dark page.
TEAMWORK_BADGE = IMAGE_DIR / "badge_teamwork.png"
SNAKE_BADGE = IMAGE_DIR / "badge_snake.png"
# Photos rather than cut-outs, so they are cropped to a ringed disc instead of keyed.
CURRY_BADGE = IMAGE_DIR / "badge_curry.png"
THINKER_BADGE = IMAGE_DIR / "badge_thinker.png"
# A heavy "B" on an editor-style button tile, drawn once in Arial Black.
BOLD_BADGE = IMAGE_DIR / "badge_bold.png"

MIN_GAMES = 5
# Scores are reported on the scale of a 5-player table.
SCORE_SCALE = 5
SCORE_SCALED = f"Average Score ({SCORE_SCALE}-player scale)"
HAND_SCALED = f"Average Hand Played ({SCORE_SCALE}-player scale)"
SNAKE_COLUMN = "Peer Survival Rate | Survived"
TEAM_COLUMN = "Peer Survival Rate | All Rounds"
ELO_START = 1000
ELO_K = 32

# Normalized columns divide by the table size: an n-player round scores in [-n, n]
# with 0 excluded, so score/n puts every table on one -1..1 scale. Stacked house-rule
# penalties can push a score past -n, and the normalized value past -1 with it.
STAT_COLUMNS = [
    "Player",
    "Games",
    # Each raw count sits beside the rate it drives.
    "Rounds Won",
    "Win Rate",
    "Solo Wins",
    "Rounds Survived",
    "Survival Rate",
    "Total Score",
    "Average Score",
    "Average Hand Played",
    "Average Group Score",
    TEAM_COLUMN,
    SNAKE_COLUMN,
    "Average Delta",
    "Elo",
    # The normalized block sits together at the far right.
    "Average Normalized Score",
    "Average Normalized Hand Played",
    "Average Normalized Group Score",
    "Average Normalized Delta",
]

DATE_RE = re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$")
END_CARDS = ("S", "X")  # syringe, toxin
SCORE_RE = re.compile(r"^([+-]?\d+|[SX])((?:[+-]\d+)*)$")
# Names and scores in the round-per-line format may be split by commas, spaces or
# both; a parenthesised adjustment like "(5 - 1)" stays one token.
TOKEN_RE = re.compile(r"\([^)]*\)|[^,\s]+")
# Misspellings and alternate spellings, keyed and valued in lower case.
NAME_ALIASES = {
    "katna": "karna",
    "samir": "sameer",
}

PAGE_W, PAGE_H = 8.5, 11.0
SURFACE = "#0f1512"
INK = "#ffffff"
INK_SOFT = "#c3c2b7"
INK_MUTED = "#898781"
GRID = "#2c2c2a"
AXIS = "#3d3d39"
LATTICE = "#243029"
GLASS = "#33463c"

# Accents validated for the dark surface: lightness band, chroma, CVD separation in
# reading order and 3:1 contrast. More vivid steps fail the dark lightness band. Blue
# is deliberately absent -- against the violet it lands at dE 9.8 for normal vision.
# Each award owns one accent and is direct-labelled, so no panel decodes by hue.
ACCENTS = ("#9085e9", "#199e70", "#c98500", "#008300", "#e66767")

SECTIONS_PER_PAGE = 3
# The badge sits in a fixed slot left of the heading; wide art is capped to fit it.
BADGE_SLOT_W = 0.62


def parse_score(token, table_size):
    """Return (score, hand) for a token like '5', '-3', '(5-1)', 'S' or '(X-1)'.

    S (syringe) and X (toxin) mark a player who ended the game with that card, which
    the house rules score as -ceil(n/2) at an n-player table. The hand is how big a
    card was gambled on: the score's size, but 0 for S or X, since that penalty is
    set by the table rather than chosen.
    """
    text = token.strip()
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1]
    text = text.replace(" ", "").upper()
    match = SCORE_RE.match(text)
    if not match:
        raise ValueError(f"Cannot parse score {token!r}")
    head = match.group(1)
    base = -((table_size + 1) // 2) if head in END_CARDS else int(head)
    adjustments = [int(a) for a in re.findall(r"[+-]\d+", match.group(2))]
    score = base + sum(adjustments)
    return score, 0 if head in END_CARDS else abs(score)


def parse_date(text):
    for fmt in ("%m/%d/%y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def canonical_name(text):
    """Fold case and known misspellings, so 'karna', 'Karna' and 'katna' are one player."""
    key = text.strip().lower()
    return NAME_ALIASES.get(key, key).capitalize()


def parse_sessions(path):
    """Read both score formats; a date block may use either.

    Per player:  'Tony: 1, 3, 1' -- one line per player, one column per round.
    Per round:   a line of names, then a line of their scores in the same order.
    Any digit-free line before a names line (e.g. 'post chai') is a label and ignored.
    """
    sessions = []
    current = None
    pending = None  # (lineno, names) awaiting their scores line
    for lineno, raw in enumerate(path.read_text().splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue
        if DATE_RE.match(line):
            current = {"date": line, "parsed_date": parse_date(line), "players": {}, "rounds": []}
            sessions.append(current)
            pending = None
            continue
        if current is None:
            print(f"Line {lineno}: scores before any date header, skipping {line!r}")
            continue

        if ":" in line:
            # Kept as text: an S or X is worth -ceil(n/2), and n isn't known until
            # every player line under this date has been read.
            name, _, scores = line.partition(":")
            tokens = [token for token in scores.split(",") if token.strip()]
            current["players"][canonical_name(name)] = (lineno, tokens)
            continue

        tokens = TOKEN_RE.findall(line)
        if not any(ch.isdigit() for ch in line):
            # A later digit-free line replaces an earlier one, so a label line is
            # simply overwritten by the names line that follows it.
            pending = (lineno, [canonical_name(token) for token in tokens])
            continue
        if pending is None:
            print(f"Line {lineno}: scores with no names line above, skipping {line!r}")
            continue
        names_line, names = pending
        pending = None
        if len(names) != len(tokens):
            print(f"Line {lineno}: {len(names)} names (line {names_line}) but "
                  f"{len(tokens)} scores -- skipping this round")
            continue
        if len(set(names)) != len(names):
            print(f"Line {names_line}: a player is listed twice -- skipping this round")
            continue
        try:
            current["rounds"].append(
                {name: parse_score(token, len(names)) for name, token in zip(names, tokens)}
            )
        except ValueError as err:
            print(f"Line {lineno}: {err} -- skipping this round")
    return sessions


def compute_elo(round_scores):
    """Pairwise Elo on clamped scores: every player in a game is rated against every other."""
    ratings = {}
    for scores in round_scores:
        names = list(scores)
        if len(names) < 2:
            continue
        # A score's sign is survival and its magnitude is only how big a card was
        # gambled on, so dying on -1 is no win over dying on -4. Clamping ties every
        # death together and leaves the ordering among survivors untouched. Done here
        # rather than at the call site so round_scores stays raw for the other stats.
        scores = {name: max(value, 0) for name, value in scores.items()}
        for name in names:
            ratings.setdefault(name, ELO_START)

        deltas = dict.fromkeys(names, 0.0)
        k = ELO_K / (len(names) - 1)
        for index, a in enumerate(names):
            for b in names[index + 1:]:
                expected = 1 / (1 + 10 ** ((ratings[b] - ratings[a]) / 400))
                actual = 0.5 if scores[a] == scores[b] else float(scores[a] > scores[b])
                deltas[a] += k * (actual - expected)
                deltas[b] -= k * (actual - expected)

        # Applied after the pair loop so every pair is rated pre-game.
        for name in names:
            ratings[name] += deltas[name]
    return ratings


def average(records, pick):
    return sum(pick(record) for record in records) / len(records)


def build_tables(sessions):
    log_rows = []
    player_order = []
    seen = set()
    played = []

    for session_index, session in enumerate(sessions):
        date, players = session["date"], session["players"]
        if session["parsed_date"] is None:
            print(f"Error: cannot read date {date!r} -- skipping this date")
            continue

        # Both formats end up as one {name: (score, hand)} dict per round, in play order.
        session_rounds = list(session["rounds"])
        if players:
            parsed = {}
            for name, (lineno, tokens) in players.items():
                games = []
                for token in tokens:
                    try:
                        games.append(parse_score(token, len(players)))
                    except ValueError as err:
                        print(f"Line {lineno}: {err}")
                parsed[name] = games
            players = parsed
            # A score that failed to parse leaves its player a round short, so it is
            # caught here along with any genuinely missing score.
            counts = {name: len(games) for name, games in players.items()}
            if len(set(counts.values())) != 1:
                print(f"Error: {date} has uneven game counts {counts} -- "
                      f"skipping its per-player lines")
            else:
                session_rounds += [
                    {name: games[round_index] for name, games in players.items()}
                    for round_index in range(next(iter(counts.values())))
                ]

        for round_index, entries in enumerate(session_rounds):
            for name in entries:
                if name not in seen:
                    seen.add(name)
                    player_order.append(name)
            scores = {name: score for name, (score, _) in entries.items()}
            hands = {name: hand for name, (_, hand) in entries.items()}
            log_rows.append({"Date": date, "Round": round_index + 1, **scores})
            played.append(((session["parsed_date"], session_index, round_index), scores, hands))

    played.sort(key=lambda entry: entry[0])
    # Only the hand-played stats need hands; everything downstream takes plain scores.
    round_scores = [scores for _, scores, _ in played]
    ratings = compute_elo(round_scores)

    # Every per-player stat below needs the round it was earned in -- the table size
    # sets the score range, and the group's own results are the yardstick.
    stats = {name: [] for name in player_order}
    for _, scores, hands in played:
        count = len(scores)
        values = list(scores.values())
        group_score = sum(values) / count
        survivors = sum(value > 0 for value in values)
        # A round is only won by the top score, and only if it survived; a table that
        # all went down together has no winner. Tied leaders each take the win.
        best = max(values)
        for name, score in scores.items():
            # Both group-survival stats (team player and snake) count everyone BUT
            # the player. Counting the player would put a floor of 1/n under every
            # round they survived, and that floor moves with the table size --
            # exactly the confound the normalized columns exist to kill.
            peers = (survivors - (score > 0)) / (count - 1) if count > 1 else None
            stats[name].append(
                {
                    "score": score,
                    "hand": hands[name],
                    "won": score > 0 and score == best,
                    # The only survivor at a table of two or more.
                    "solo": score > 0 and survivors == 1 and count > 1,
                    "count": count,
                    "group_score": group_score,
                    "peer_survival": peers,
                    "delta": score - group_score,
                }
            )

    game_log = pd.DataFrame(log_rows, columns=["Date", "Round", *player_order])
    for name in player_order:
        game_log[name] = game_log[name].astype("Int64")

    stat_rows = []
    for name in player_order:
        games = stats[name]
        with_peers = [g for g in games if g["peer_survival"] is not None]
        survived = [g for g in with_peers if g["score"] > 0]
        stat_rows.append(
            {
                "Player": name,
                "Games": len(games),
                "Rounds Won": sum(g["won"] for g in games),
                "Win Rate": round(average(games, lambda g: g["won"]), 3),
                "Solo Wins": sum(g["solo"] for g in games),
                "Rounds Survived": sum(g["score"] > 0 for g in games),
                "Survival Rate": round(average(games, lambda g: g["score"] > 0), 3),
                "Total Score": sum(g["score"] for g in games),
                "Average Score": round(average(games, lambda g: g["score"]), 2),
                # Normalized stats divide EACH round by its own table size and then
                # average those per-round ratios -- never total/total, which would let
                # the biggest tables dominate.
                "Average Normalized Score":
                    round(average(games, lambda g: g["score"] / g["count"]), 3),
                "Average Hand Played": round(average(games, lambda g: g["hand"]), 2),
                "Average Normalized Hand Played":
                    round(average(games, lambda g: g["hand"] / g["count"]), 3),
                "Average Group Score": round(average(games, lambda g: g["group_score"]), 2),
                "Average Normalized Group Score":
                    round(average(games, lambda g: g["group_score"] / g["count"]), 3),
                TEAM_COLUMN:
                    round(average(with_peers, lambda g: g["peer_survival"]), 3)
                    if with_peers else None,
                # Blank, not 0, when a player never survived a round: "never snaked"
                # and "never had the chance" are different claims.
                SNAKE_COLUMN:
                    round(average(survived, lambda g: g["peer_survival"]), 3)
                    if survived else None,
                "Average Delta": round(average(games, lambda g: g["delta"]), 2),
                "Average Normalized Delta":
                    round(average(games, lambda g: g["delta"] / g["count"]), 3),
                "Elo": round(ratings.get(name, ELO_START)),
            }
        )
    player_stats = pd.DataFrame(stat_rows, columns=STAT_COLUMNS)
    return game_log, player_stats, round_scores


def summarize_rounds(round_scores):
    """Pool every player-round in the bucket: one survival rate, one average score."""
    values = [score for scores in round_scores for score in scores.values()]
    return {
        "Rounds": len(round_scores),
        "Survival Rate": round(sum(v > 0 for v in values) / len(values), 3),
        "Average Score": round(sum(values) / len(values), 2),
    }


def build_group_tables(rounds):
    """Two views of the same rounds: per lineup, and per table size."""
    by_group = {}
    by_count = {}
    for scores in rounds:
        lineup = tuple(sorted(scores))
        by_group.setdefault(lineup, []).append(scores)
        by_count.setdefault(len(lineup), []).append(scores)

    group_rows = [
        {
            "Players": ", ".join(lineup),
            "Player Count": len(lineup),
            **summarize_rounds(round_scores),
        }
        for lineup, round_scores in by_group.items()
    ]
    group_rows.sort(key=lambda row: (row["Player Count"], -row["Rounds"], row["Players"]))

    count_rows = [
        {
            "Player Count": count,
            # A count backed by one lineup says as much about those players as it
            # does about the table size, so the lineup tally travels with the row.
            "Lineups": len({tuple(sorted(scores)) for scores in round_scores}),
            **summarize_rounds(round_scores),
        }
        for count, round_scores in sorted(by_count.items())
    ]

    group_stats = pd.DataFrame(
        group_rows, columns=["Players", "Player Count", "Rounds", "Survival Rate", "Average Score"]
    )
    count_stats = pd.DataFrame(
        count_rows, columns=["Player Count", "Lineups", "Rounds", "Survival Rate", "Average Score"]
    )
    return group_stats, count_stats


def pick_font(*candidates):
    available = {f.name for f in font_manager.fontManager.ttflist}
    for name in candidates:
        if name in available:
            return name
    return "DejaVu Sans"


def spaced(text, gap=2):
    return (" " * gap).join(text)


def hex_lattice(ax, x, y, rings, radius, color):
    """Faint benzene-ring lattice, used only as page texture behind the margins."""
    step_x = radius * 1.5
    step_y = radius * (3 ** 0.5)
    for row in range(rings):
        for col in range(rings):
            cx = x + col * step_x
            cy = y + row * step_y + (step_y / 2 if col % 2 else 0)
            ax.add_patch(
                RegularPolygon(
                    (cx, cy),
                    6,
                    radius=radius,
                    orientation=0,
                    facecolor="none",
                    edgecolor=color,
                    linewidth=0.5,
                    alpha=0.5,
                    zorder=0,
                )
            )


def load_badge(path):
    if not path.exists():
        print(f"{path.name} not found -- that badge is left blank")
        return None
    return Image.open(path).convert("RGBA")


def place_image(fig, page, image, center, height, max_width=None):
    """Scale to `height`, but never wider than `max_width`.

    Both are in the same nominal units: matplotlib multiplies an OffsetImage zoom by
    dpi/72 at draw time, so a badge lands about 1.39x these numbers in inches. The
    existing badge sizes are tuned to that convention, so it is kept rather than fixed.
    """
    if image is None:
        return
    array = np.asarray(image)
    zoom = height * fig.dpi / array.shape[0]
    if max_width:
        zoom = min(zoom, max_width * fig.dpi / array.shape[1])
    box = AnnotationBbox(OffsetImage(array, zoom=zoom), center, frameon=False, pad=0)
    box.set_zorder(6)
    page.add_artist(box)


def _place(page, parts, x, y, tilt, zorder=1):
    spin = Affine2D().rotate_deg_around(x, y, tilt) if tilt else None
    for part in parts:
        part.set_zorder(zorder)
        page.add_patch(part)
        if spin:
            part.set_transform(spin + page.transData)


def conical_flask(page, x, y, size, edge, liquid=None, alpha=0.9, tilt=0):
    half, neck = 0.40 * size, 0.11 * size
    body, throat = 0.52 * size, 0.26 * size
    shell = [(-half, 0), (half, 0), (neck, body), (neck, body + throat),
             (-neck, body + throat), (-neck, body)]
    parts = []
    if liquid:
        level = 0.26 * size
        wide = neck + (half - neck) * (1 - level / body)
        parts.append(Polygon([(x - half, y), (x + half, y), (x + wide, y + level),
                              (x - wide, y + level)], closed=True, facecolor=liquid,
                             edgecolor="none", alpha=alpha * 0.5))
    parts.append(Polygon([(x + dx, y + dy) for dx, dy in shell], closed=True,
                         facecolor="none", edgecolor=edge, linewidth=0.9, alpha=alpha))
    parts.append(Rectangle((x - neck * 1.7, y + body + throat), neck * 3.4, 0.03 * size,
                           facecolor=edge, edgecolor="none", alpha=alpha))
    _place(page, parts, x, y, tilt)


def test_tube(page, x, y, size, edge, liquid=None, alpha=0.9, tilt=0):
    half, body = 0.15 * size, 0.62 * size
    arc = [(x + half * np.cos(a), y + half + half * np.sin(a))
           for a in np.linspace(np.pi, 2 * np.pi, 20)]
    parts = []
    if liquid:
        top = y + half + body * 0.40
        parts.append(Polygon(arc + [(x + half, top), (x - half, top)], closed=True,
                             facecolor=liquid, edgecolor="none", alpha=alpha * 0.5))
    parts.append(Polygon([(x - half, y + half + body)] + arc + [(x + half, y + half + body)],
                         closed=True, facecolor="none", edgecolor=edge, linewidth=0.9, alpha=alpha))
    parts.append(Rectangle((x - half * 1.3, y + half + body), half * 2.6, 0.025 * size,
                           facecolor=edge, edgecolor="none", alpha=alpha))
    _place(page, parts, x, y, tilt)


def round_flask(page, x, y, size, edge, liquid=None, alpha=0.9, tilt=0):
    radius, neck_w, neck_h = 0.30 * size, 0.085 * size, 0.32 * size
    parts = []
    if liquid:
        parts.append(Wedge((x, y + radius), radius * 0.94, 200, 340, facecolor=liquid,
                           edgecolor="none", alpha=alpha * 0.5))
    parts.append(Circle((x, y + radius), radius, facecolor="none", edgecolor=edge,
                        linewidth=0.9, alpha=alpha))
    parts.append(Rectangle((x - neck_w, y + radius * 1.72), neck_w * 2, neck_h,
                           facecolor="none", edgecolor=edge, linewidth=0.9, alpha=alpha))
    parts.append(Rectangle((x - neck_w * 1.7, y + radius * 1.72 + neck_h), neck_w * 3.4,
                           0.03 * size, facecolor=edge, edgecolor="none", alpha=alpha))
    _place(page, parts, x, y, tilt)


def vial(page, x, y, size, edge, liquid=None, alpha=0.9, tilt=0):
    half, body, cap = 0.17 * size, 0.46 * size, 0.11 * size
    parts = []
    if liquid:
        parts.append(Rectangle((x - half, y), half * 2, body * 0.44, facecolor=liquid,
                               edgecolor="none", alpha=alpha * 0.5))
    parts.append(Rectangle((x - half, y), half * 2, body, facecolor="none", edgecolor=edge,
                           linewidth=0.9, alpha=alpha))
    parts.append(Rectangle((x - half * 0.62, y + body), half * 1.24, cap, facecolor=edge,
                           edgecolor="none", alpha=alpha * 0.85))
    _place(page, parts, x, y, tilt)


def syringe(page, x, y, size, edge, alpha=0.9, tilt=0):
    barrel, bore = 0.58 * size, 0.10 * size
    parts = [
        Rectangle((x, y - bore), barrel, bore * 2, facecolor="none", edgecolor=edge,
                  linewidth=0.9, alpha=alpha),
        Rectangle((x + barrel, y - bore * 0.22), 0.30 * size, bore * 0.44, facecolor=edge,
                  edgecolor="none", alpha=alpha),
        Rectangle((x - 0.09 * size, y - bore * 1.6), 0.09 * size, bore * 3.2, facecolor=edge,
                  edgecolor="none", alpha=alpha * 0.85),
        Rectangle((x + barrel * 0.18, y - bore * 0.8), barrel * 0.45, bore * 1.6,
                  facecolor=edge, edgecolor="none", alpha=alpha * 0.35),
    ]
    _place(page, parts, x, y, tilt)


def benzene(page, x, y, size, edge, alpha=0.9):
    parts = [
        RegularPolygon((x, y), 6, radius=size, orientation=0, facecolor="none",
                       edgecolor=edge, linewidth=0.9, alpha=alpha),
        Circle((x, y), size * 0.55, facecolor="none", edgecolor=edge, linewidth=0.7,
               alpha=alpha * 0.8),
    ]
    for angle in (0, 120, 240):
        parts.append(Circle((x + size * 1.42 * np.cos(np.radians(angle)),
                             y + size * 1.42 * np.sin(np.radians(angle))),
                            size * 0.17, facecolor=edge, edgecolor="none", alpha=alpha))
    _place(page, parts, x, y, 0)


def bubbles(page, x, y, size, edge, alpha=0.9):
    for dx, dy, scale in ((0, 0, 1.0), (0.42, 0.30, 0.66), (0.18, 0.62, 0.45),
                          (0.62, 0.72, 0.34)):
        page.add_patch(Circle((x + dx * size, y + dy * size), size * 0.17 * scale,
                              facecolor="none", edgecolor=edge, linewidth=0.7,
                              alpha=alpha, zorder=1))


def draw_ranking(fig, bounds, frame, section, accent, fonts):
    """One ranked panel. Form is chosen by what the numbers mean, not by taste."""
    column, form, fmt = section["column"], section["form"], section["value_fmt"]
    ax = fig.add_axes([bounds[0] / PAGE_W, bounds[1] / PAGE_H, bounds[2] / PAGE_W, bounds[3] / PAGE_H])
    ax.set_facecolor("none")
    ax.set_zorder(2)

    # A section may rank upward (the snake award: least peer survival wins).
    ranked = (frame.dropna(subset=[column])
              .sort_values(column, ascending=section.get("ascending", False))
              .reset_index(drop=True))
    if ranked.empty:
        ax.axis("off")
        ax.text(0.0, 0.5, "not enough rounds yet", transform=ax.transAxes, va="center",
                fontname=fonts["body"], fontsize=8.5, color=INK_MUTED)
        return
    values = ranked[column].tolist()
    positions = list(range(len(values)))

    low, high = min(values), max(values)
    if form == "dot":
        # Elo has no meaningful zero, so a bar from 0 would be dishonest: dots on a
        # shared axis compare the ratings without implying a baseline.
        pad = max((high - low) * 0.22, 6)
        left, right = low - pad, high + pad
        baseline = None
    elif form == "bar":
        # A rate cannot exceed 100%, so the axis stops there; the slack past it is
        # label gutter, kept tick-free so it doesn't read as part of the scale.
        left, right = 0.0, min(high * 1.42, 1.18)
        baseline = 0.0
    elif form == "count":
        # A tally has a true zero and no ceiling; the slack is label gutter again.
        left, right = 0.0, max(high * 1.35, high + 1)
        baseline = 0.0
        ax.xaxis.set_major_locator(MaxNLocator(integer=True))
    else:
        # Average score is a signed quantity, so it is read against a visible zero.
        span = max(abs(low), abs(high))
        left, right = low - span * 0.34, max(high, 0) + span * 0.22
        baseline = 0.0

    ax.set_xlim(left, right)
    ax.set_ylim(len(values) - 0.45, -0.75)
    ax.grid(axis="x", color=GRID, linewidth=0.55, zorder=0)
    ax.set_axisbelow(True)
    ax.tick_params(axis="x", colors=INK_MUTED, labelsize=6.5, length=0, pad=5)
    ax.set_yticks([])
    ax.xaxis.set_major_formatter(FuncFormatter(section["tick_fmt"]))
    for spine in ax.spines.values():
        spine.set_visible(False)
    for label in ax.get_xticklabels():
        label.set_fontname(fonts["num"])

    if baseline is not None:
        ax.axvline(baseline, color=AXIS, linewidth=0.9, zorder=1)

    thickness = 8.5
    for position, value in zip(positions, values):
        if form == "dot":
            ax.plot([left, value], [position, position], color=GRID, linewidth=0.7,
                    linestyle=(0, (1, 2.6)), zorder=1)
            ax.plot([value], [position], marker="o", markersize=9.5, color=accent,
                    markeredgecolor=SURFACE, markeredgewidth=1.4, zorder=3)
        else:
            # Butt-capped line plus a round marker at the data end: the rounded end
            # lands on the value, the baseline end stays square against the axis.
            ax.plot([baseline, value], [position, position], color=accent,
                    linewidth=thickness, solid_capstyle="butt", zorder=3)
            ax.plot([value], [position], marker="o", markersize=thickness,
                    color=accent, markeredgecolor="none", zorder=3)

    span = right - left
    offset = span * 0.022
    for position, (value, name) in enumerate(zip(values, ranked["Player"])):
        leader = position == 0
        if form == "diverge" and value < 0:
            text_x, align = value - offset, "right"
        else:
            text_x, align = value + offset, "left"
        ax.text(text_x, position, fmt(value), ha=align, va="center",
                fontname=fonts["num"], fontsize=8.6,
                color=INK if leader else INK_SOFT, zorder=4)

        transform = ax.get_yaxis_transform()
        ax.text(-0.175, position, str(position + 1), ha="left", va="center",
                transform=transform, fontname=fonts["num"], fontsize=7.5,
                color=INK if leader else INK_MUTED, zorder=4)
        ax.text(-0.028, position, name, ha="right", va="center", transform=transform,
                fontname=fonts["body"], fontsize=9.4,
                color=INK if leader else INK_SOFT,
                fontweight="demi" if leader else "normal", zorder=4)


def decorate_section(page, index, top, accent):
    """Glassware in the empty strip right of each heading, tinted to the section.

    Each section gets its own arrangement so the three strips don't read as one
    repeated stamp.
    """
    base = top - 0.58
    tint = dict(liquid=accent, alpha=0.32)
    index %= 3
    if index == 0:
        conical_flask(page, 5.94, base, 0.62, accent, tilt=-7, **tint)
        test_tube(page, 6.44, base + 0.03, 0.54, accent, tilt=12, **tint)
        benzene(page, 6.84, base + 0.22, 0.12, accent, alpha=0.34)
        vial(page, 7.22, base + 0.02, 0.50, accent, tilt=-4, **tint)
        bubbles(page, 7.52, base + 0.12, 0.34, accent, alpha=0.30)
    elif index == 1:
        round_flask(page, 5.98, base, 0.60, accent, **tint)
        test_tube(page, 6.42, base + 0.02, 0.50, accent, tilt=-14, **tint)
        test_tube(page, 6.70, base + 0.04, 0.56, accent, tilt=9, **tint)
        bubbles(page, 7.02, base + 0.30, 0.36, accent, alpha=0.30)
        conical_flask(page, 7.42, base, 0.56, accent, tilt=5, **tint)
        benzene(page, 7.76, base + 0.26, 0.11, accent, alpha=0.34)
    else:
        syringe(page, 5.82, base + 0.22, 0.70, accent, alpha=0.34, tilt=-11)
        vial(page, 6.62, base + 0.01, 0.52, accent, tilt=6, **tint)
        conical_flask(page, 7.02, base, 0.58, accent, tilt=-5, **tint)
        test_tube(page, 7.46, base + 0.03, 0.52, accent, tilt=15, **tint)
        benzene(page, 7.76, base + 0.30, 0.12, accent, alpha=0.34)


def decorate_margins(page, fonts):
    """Page texture: ring lattices, a spread of glassware, and stray formulae."""
    hex_lattice(page, 6.9, 9.78, 4, 0.30, LATTICE)
    hex_lattice(page, -0.55, 0.15, 3, 0.30, LATTICE)

    # Top-right, beside the title block.
    conical_flask(page, 6.02, 9.58, 0.74, GLASS, liquid=GLASS, tilt=-9)
    test_tube(page, 6.52, 9.66, 0.62, GLASS, liquid=GLASS, tilt=13)
    syringe(page, 6.85, 10.26, 0.62, GLASS, alpha=0.75, tilt=-8)
    round_flask(page, 7.92, 9.98, 0.60, GLASS, liquid=GLASS)
    benzene(page, 7.52, 9.53, 0.13, GLASS)

    # Bottom strip, under the footer rule.
    conical_flask(page, 6.28, 0.16, 0.56, GLASS, liquid=GLASS, tilt=6)
    test_tube(page, 6.76, 0.18, 0.50, GLASS, liquid=GLASS, tilt=-13)
    round_flask(page, 7.22, 0.16, 0.52, GLASS, liquid=GLASS)
    vial(page, 7.60, 0.18, 0.46, GLASS, liquid=GLASS, tilt=8)
    benzene(page, 5.95, 0.34, 0.11, GLASS)
    test_tube(page, 0.34, 0.16, 0.44, GLASS, liquid=GLASS, tilt=9)
    bubbles(page, 5.55, 0.22, 0.30, GLASS)

    vial(page, 0.35, 7.55, 0.42, GLASS, liquid=GLASS, tilt=-7)
    benzene(page, 0.36, 4.95, 0.10, GLASS)
    conical_flask(page, 0.35, 1.95, 0.44, GLASS, liquid=GLASS, tilt=6)

    for y, formula in ((8.05, "C₈H₁₁NO₂"), (6.05, "C₁₇H₁₉NO₃"), (3.35, "NaCN · H₂SO₄")):
        page.text(0.37, y, formula, fontname=fonts["num"], fontsize=6.2, color="#3c5147",
                  rotation=90, va="bottom", ha="center", zorder=1)


def render_page(page_sections, first_index, qualified, fonts, subtitle, page_no, page_count):
    """One sheet: masthead, up to SECTIONS_PER_PAGE awards, footer."""
    fig = plt.figure(figsize=(PAGE_W, PAGE_H))
    fig.patch.set_facecolor(SURFACE)

    page = fig.add_axes([0, 0, 1, 1])
    page.set_xlim(0, PAGE_W)
    page.set_ylim(0, PAGE_H)
    page.set_aspect("equal")
    page.axis("off")
    page.set_facecolor(SURFACE)
    page.set_zorder(1)

    decorate_margins(page, fonts)

    page.plot([0.7, 7.8], [10.55, 10.55], color=AXIS, linewidth=1.1, zorder=2)
    page.plot([0.7, 7.8], [10.50, 10.50], color=AXIS, linewidth=0.5, zorder=2)
    page.text(0.7, 10.24, spaced("CLINICAL TOXICOLOGY · PLAYER ASSESSMENT"),
              fontname=fonts["label"], fontsize=6.8, color=INK_MUTED, zorder=3)
    page.text(0.7, 9.72, "Antidote Report Card", fontname=fonts["title"], fontsize=34,
              color=INK, zorder=3)
    page.text(0.72, 9.44, subtitle if page_no == 1 else f"{subtitle} · continued",
              fontname=fonts["body"], fontsize=9, color=INK_SOFT, zorder=3)
    page.text(7.8, 9.44, f"Minimum {MIN_GAMES} rounds to qualify", ha="right",
              fontname=fonts["body"], fontsize=9, color=INK_SOFT, zorder=3)
    page.plot([0.7, 7.8], [9.24, 9.24], color=AXIS, linewidth=0.6, zorder=2)

    # Accent and glassware follow the award's place in the whole report; the vertical
    # slot follows its place on this sheet.
    for offset, section in enumerate(page_sections):
        index = first_index + offset
        # More awards than accents: the cycle still keeps every accent on a sheet distinct.
        accent = ACCENTS[index % len(ACCENTS)]
        top = 9.05 - offset * 2.78

        place_image(fig, page, section["badge"], (1.06, top - 0.33), section["badge_height"],
                    BADGE_SLOT_W)
        page.text(1.58, top - 0.24, spaced(section["heading"]), fontname=fonts["label"],
                  fontsize=11, color=INK, zorder=3)
        page.text(1.60, top - 0.52, section["blurb"], fontname=fonts["body"], fontsize=8,
                  color=INK_MUTED, zorder=3)
        decorate_section(page, index, top, accent)

        draw_ranking(fig, (1.75, top - 2.46, 6.05, 1.62), qualified, section, accent, fonts)

    page.plot([0.7, 7.8], [0.62, 0.62], color=AXIS, linewidth=0.6, zorder=2)
    page.text(0.7, 0.40,
              f"Minimum {MIN_GAMES} rounds to qualify · scores normalized to a "
              f"{SCORE_SCALE}-player table · generated {datetime.now():%d %B %Y}",
              fontname=fonts["body"], fontsize=6.8, color=INK_MUTED, zorder=3)
    page.text(7.8, 0.40, f"Page {page_no} of {page_count}", ha="right",
              fontname=fonts["body"], fontsize=6.8, color=INK_MUTED, zorder=3)
    return fig


def write_report_card(player_stats, game_log):
    # PlayerStats keeps everyone; the rankings only mean something once a player has
    # enough rounds for the rates to settle, so the report card qualifies them here.
    qualified = player_stats[player_stats["Games"] >= MIN_GAMES].reset_index(drop=True)
    for name, games in zip(player_stats["Player"], player_stats["Games"]):
        if games < MIN_GAMES:
            print(f"{name} has {games} rounds (min {MIN_GAMES}) -- not ranked on the report card")
    if qualified.empty:
        print(f"No player has {MIN_GAMES} rounds yet -- skipping the report card")
        return

    qualified = qualified.assign(**{
        # Put every table on the scale of a 5-player game, so a -5 at a five-player
        # table and a -3 at a three-player one are read as the same performance.
        SCORE_SCALED: qualified["Average Normalized Score"] * SCORE_SCALE,
        # Card values run up to the table size too, so hands get the same treatment.
        HAND_SCALED: qualified["Average Normalized Hand Played"] * SCORE_SCALE,
    })

    fonts = {
        "title": pick_font("Baskerville", "Hoefler Text", "Didot", "Georgia", "DejaVu Serif"),
        "label": pick_font("Copperplate", "Optima", "Futura", "DejaVu Sans"),
        "body": pick_font("Avenir Next", "Avenir", "Helvetica Neue", "DejaVu Sans"),
        "num": pick_font("Menlo", "Courier New", "DejaVu Sans Mono"),
    }

    muscle = load_badge(MUSCLE_BADGE)
    survivor = load_badge(SURVIVOR_LOGO)
    trophy = load_badge(TROPHY_BADGE)
    teamwork = load_badge(TEAMWORK_BADGE)
    snake = load_badge(SNAKE_BADGE)
    curry = load_badge(CURRY_BADGE)
    thinker = load_badge(THINKER_BADGE)
    bold = load_badge(BOLD_BADGE)

    sections = [
        {
            "badge": curry,
            "badge_height": 0.48,
            "heading": "BORN TO WIN",
            "blurb": "Share of rounds finished with the top positive score · ties share it",
            "column": "Win Rate",
            "form": "bar",
            "value_fmt": lambda v: f"{v * 100:.1f}%",
            "tick_fmt": lambda v, _: f"{v * 100:.0f}%",
        },
        {
            "badge": survivor,
            "badge_height": 0.46,
            "heading": "BEST SURVIVAL INSTINCTS",
            "blurb": "Share of rounds finished on a positive score",
            "column": "Survival Rate",
            "form": "bar",
            "value_fmt": lambda v: f"{v * 100:.1f}%",
            "tick_fmt": lambda v, _: f"{v * 100:.0f}%",
        },
        {
            "badge": trophy,
            "badge_height": 0.50,
            "heading": "SCORING CHAMPION",
            "blurb": f"Average score per round · normalized to a {SCORE_SCALE}-player table",
            "column": SCORE_SCALED,
            "form": "diverge",
            # Adding 0.0 after rounding turns float noise like -1e-17 into 0.00, not -0.00.
            "value_fmt": lambda v: f"{round(v, 2) + 0.0:.2f}",
            "tick_fmt": lambda v, _: f"{v:.0f}",
        },
        {
            "badge": muscle,
            "badge_height": 0.50,
            "heading": "ELO POWER RANKINGS",
            "blurb": "Pairwise head-to-head rating · everyone starts at 1000",
            "column": "Elo",
            "form": "dot",
            "value_fmt": lambda v: f"{v:,.0f}",
            "tick_fmt": lambda v, _: f"{v:,.0f}",
        },
        {
            # The teamwork art is far wider than it is tall, so it is set shorter to
            # keep clear of the heading.
            "badge": teamwork,
            "badge_height": 0.34,
            "heading": "BEST TEAM PLAYER",
            "blurb": "Share of the rest of the table that survived · all rounds",
            "column": TEAM_COLUMN,
            "form": "bar",
            "value_fmt": lambda v: f"{v * 100:.1f}%",
            "tick_fmt": lambda v, _: f"{v * 100:.0f}%",
        },
        {
            "badge": snake,
            "badge_height": 0.50,
            "heading": "BIGGEST SNAKE",
            "blurb": "Share of the rest of the table that survived when they did · lowest first",
            "column": SNAKE_COLUMN,
            "form": "bar",
            "ascending": True,
            "value_fmt": lambda v: f"{v * 100:.1f}%",
            "tick_fmt": lambda v, _: f"{v * 100:.0f}%",
        },
        {
            "badge": thinker,
            "badge_height": 0.48,
            "heading": "FREE THINKER",
            "blurb": "Solo wins · rounds they alone survived while the rest of the table died",
            "column": "Solo Wins",
            "form": "count",
            "value_fmt": lambda v: f"{v:.0f}",
            "tick_fmt": lambda v, _: f"{v:.0f}",
        },
        {
            "badge": bold,
            "badge_height": 0.46,
            "heading": "BOLDEST PLAYER",
            "blurb": f"Average size of the hand played, win or lose · syringe and toxin count 0 "
                     f"· {SCORE_SCALE}-player scale",
            "column": HAND_SCALED,
            # A magnitude with a true zero, read like a tally rather than a rate.
            "form": "count",
            "value_fmt": lambda v: f"{v:.2f}",
            "tick_fmt": lambda v, _: f"{v:.0f}",
        },
    ]

    dates = sorted({d for d in game_log["Date"]}, key=lambda d: parse_date(d) or datetime.min.date())
    latest = parse_date(dates[-1]) if dates else None
    rounds_word = "round" if len(game_log) == 1 else "rounds"
    sessions_word = "session" if len(dates) == 1 else "sessions"
    subtitle = f"{len(game_log)} {rounds_word} · {len(dates)} {sessions_word}"
    if latest:
        subtitle += f" · through {latest.strftime('%d %B %Y').lstrip('0')}"

    pages = [sections[start:start + SECTIONS_PER_PAGE]
             for start in range(0, len(sections), SECTIONS_PER_PAGE)]
    with PdfPages(REPORT_PDF) as pdf:
        for page_no, page_sections in enumerate(pages, start=1):
            fig = render_page(page_sections, (page_no - 1) * SECTIONS_PER_PAGE, qualified,
                              fonts, subtitle, page_no, len(pages))
            pdf.savefig(fig, facecolor=SURFACE)
            plt.close(fig)
    return len(pages), len(sections)


def cell_value(value):
    if value is None or pd.isna(value):
        return None
    return value.item() if hasattr(value, "item") else value


def stored_widths(sheet):
    """Map header text -> hand-set width, so widths follow a column when it moves.

    Excel stores widths against the letter. Keying them to the heading instead means
    a column that is reordered -- or pushed right by an inserted one -- carries the
    width you gave it rather than inheriting whatever now sits in its old slot.
    """
    widths = {}
    if sheet.max_row < 1:
        return widths
    for index, cell in enumerate(sheet[1], start=1):
        dimension = sheet.column_dimensions.get(get_column_letter(index))
        if cell.value is not None and dimension is not None and dimension.width:
            widths[str(cell.value)] = dimension.width
    return widths


def write_sheet(book, name, frame):
    if name in book.sheetnames:
        sheet = book[name]
        kept = stored_widths(sheet)
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet = book.create_sheet(name)
        kept = {}

    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        sheet.append([cell_value(value) for value in row])

    for index, column in enumerate(frame.columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = kept.get(column, max(len(str(column)) + 2, 10))

    # A dropped column would otherwise leave its width behind on the trailing letter.
    for index in range(len(frame.columns) + 1, sheet.max_column + 2):
        letter = get_column_letter(index)
        if letter in sheet.column_dimensions:
            del sheet.column_dimensions[letter]


def write_blocks(book, name, blocks):
    """Stack several titled tables down one sheet, separated by a blank row."""
    if name in book.sheetnames:
        sheet = book[name]
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet = book.create_sheet(name)

    widths = {}
    for position, (title, frame) in enumerate(blocks):
        if position:
            sheet.append([])
        sheet.append([title])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False):
            sheet.append([cell_value(value) for value in row])

        for index, column in enumerate(frame.columns, start=1):
            longest = max((len(str(value)) for value in frame.iloc[:, index - 1]), default=0)
            widths[index] = max(widths.get(index, 10), len(str(column)) + 2, min(longest + 2, 46))

    # Only size columns Excel has no stored width for, so manual widths survive.
    for index, width in widths.items():
        letter = get_column_letter(index)
        if letter in sheet.column_dimensions and sheet.column_dimensions[letter].width:
            continue
        sheet.column_dimensions[letter].width = width


def main():
    sessions = parse_sessions(INPUT_FILE)
    game_log, player_stats, rounds = build_tables(sessions)
    group_stats, count_stats = build_group_tables(rounds)

    if WORKBOOK.exists():
        book = load_workbook(WORKBOOK)
    else:
        book = Workbook()
        book.remove(book.active)

    write_sheet(book, "GameLog", game_log)
    write_sheet(book, "PlayerStats", player_stats)
    write_blocks(book, "Group Stats", [
        ("By lineup", group_stats),
        ("By player count", count_stats),
    ])
    book.save(WORKBOOK)
    report = write_report_card(player_stats, game_log)

    # Excel keeps its own copy in memory and would overwrite this save on close.
    if (BASE_DIR / f"~${WORKBOOK.name}").exists():
        print(f"Warning: {WORKBOOK.name} looks open in Excel -- close without saving, then reopen")

    ranked = int((player_stats["Games"] >= MIN_GAMES).sum())
    print(f"{WORKBOOK.name}: {len(game_log)} games, {len(player_stats)} players in stats "
          f"({ranked} ranked), {len(group_stats)} lineups across "
          f"{len(count_stats)} table sizes")
    if report:
        pages, awards = report
        print(f"{REPORT_PDF.name}: {awards} awards over {pages} pages")


if __name__ == "__main__":
    main()
